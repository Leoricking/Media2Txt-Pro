import json
import os
from typing import Dict, Any, Optional


def build_export_payload(
    title: str,
    transcript: str,
    summary: str,
    metadata: Optional[Dict[str, Any]] = None,
    raw_text: str = "",
) -> Dict[str, Any]:
    return {
        "title": (title or "").strip(),
        "transcript": transcript or "",
        "summary": summary or "",
        "metadata": dict(metadata or {}),
        "raw_text": raw_text or "",
    }


def _ensure_parent_dir(path: str):
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)


def _render_txt(payload: Dict[str, Any]) -> str:
    title = payload.get("title", "").strip()
    transcript = payload.get("transcript", "")
    summary = payload.get("summary", "")
    metadata = payload.get("metadata", {}) or {}
    raw_text = payload.get("raw_text", "")

    if transcript or summary:
        lines = []
        if title:
            lines.append(f"=== {title} ===")
            lines.append("")

        if metadata:
            for k, v in metadata.items():
                lines.append(f"[{k}] {v}")
            lines.append("")

        if transcript:
            lines.append("[Transcript]")
            lines.append(transcript)
            lines.append("")

        if summary:
            lines.append("[Summary]")
            lines.append(summary)
            lines.append("")

        return "\n".join(lines).strip()

    return raw_text.strip()


def _render_md(payload: Dict[str, Any]) -> str:
    title = payload.get("title", "").strip() or "Untitled"
    transcript = payload.get("transcript", "")
    summary = payload.get("summary", "")
    metadata = payload.get("metadata", {}) or {}
    raw_text = payload.get("raw_text", "")

    lines = [f"# {title}", ""]

    if metadata:
        lines.append("## Metadata")
        lines.append("")
        for k, v in metadata.items():
            lines.append(f"- **{k}**: {v}")
        lines.append("")

    if transcript:
        lines.append("## Transcript")
        lines.append("")
        lines.append(transcript)
        lines.append("")

    if summary:
        lines.append("## Summary")
        lines.append("")
        lines.append(summary)
        lines.append("")

    if not transcript and not summary and raw_text:
        lines.append("## Content")
        lines.append("")
        lines.append(raw_text)
        lines.append("")

    return "\n".join(lines).strip()


def _render_json(payload: Dict[str, Any]) -> str:
    data = {
        "title": payload.get("title", "").strip(),
        "metadata": payload.get("metadata", {}) or {},
        "transcript": payload.get("transcript", ""),
        "summary": payload.get("summary", ""),
        "raw_text": payload.get("raw_text", ""),
    }
    return json.dumps(data, ensure_ascii=False, indent=2)


def export_txt(path: str, payload: Dict[str, Any]) -> str:
    _ensure_parent_dir(path)
    with open(path, "w", encoding="utf-8") as f:
        f.write(_render_txt(payload))
    return path


def export_md(path: str, payload: Dict[str, Any]) -> str:
    _ensure_parent_dir(path)
    with open(path, "w", encoding="utf-8") as f:
        f.write(_render_md(payload))
    return path


def export_json(path: str, payload: Dict[str, Any]) -> str:
    _ensure_parent_dir(path)
    with open(path, "w", encoding="utf-8") as f:
        f.write(_render_json(payload))
    return path


def export_docx(path: str, payload: Dict[str, Any]) -> str:
    try:
        from docx import Document
    except Exception as e:
        raise RuntimeError(
            "python-docx is required for DOCX export.\nPlease run:\npip install python-docx"
        ) from e

    _ensure_parent_dir(path)

    title = payload.get("title", "").strip() or "Untitled"
    transcript = payload.get("transcript", "")
    summary = payload.get("summary", "")
    metadata = payload.get("metadata", {}) or {}
    raw_text = payload.get("raw_text", "")

    doc = Document()
    doc.add_heading(title, level=1)

    if metadata:
        doc.add_heading("Metadata", level=2)
        for k, v in metadata.items():
            doc.add_paragraph(f"{k}: {v}")

    if transcript:
        doc.add_heading("Transcript", level=2)
        doc.add_paragraph(transcript)

    if summary:
        doc.add_heading("Summary", level=2)
        doc.add_paragraph(summary)

    if not transcript and not summary and raw_text:
        doc.add_heading("Content", level=2)
        doc.add_paragraph(raw_text)

    doc.save(path)
    return path


def export_by_extension(path: str, payload: Dict[str, Any]) -> str:
    ext = os.path.splitext(path)[1].lower()

    if ext == ".txt":
        return export_txt(path, payload)
    if ext == ".md":
        return export_md(path, payload)
    if ext == ".json":
        return export_json(path, payload)
    if ext == ".docx":
        return export_docx(path, payload)

    # fallback
    return export_txt(path, payload)


def export_note(path: str, payload: Dict[str, Any]) -> str:
    return export_by_extension(path, payload)
